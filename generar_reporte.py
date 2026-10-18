"""
Genera el informe de precios en el FORMATO ENCOMENDADO (igual a Precios__1.xlsx):
una hoja por categoría con  ZONA · DISTRIBUIDOR · Marca · Producto · Precio entrada ·
Precio salida · Margen % · Margen Abs · Price Index · Dinamica de Sell Out  (+ N° PDV).

Fuente: Registro_Precios + maestros (Google Sheet si hay secrets.toml; si no, carpeta data/).

Uso (con valesK activo, desde la carpeta de la app):
    python generar_reporte.py                                   -> todo el histórico
    python generar_reporte.py --desde 2026-10-01 --hasta 2026-10-31

Reglas de cálculo (ajustables abajo):
    - Precios = PROMEDIO de los PDV levantados en la zona y período (sin contar "Sin existencia").
    - Margen Abs = Precio salida - Precio entrada  (margen del PDV).
    - Margen %   = Margen Abs / Precio salida      (margen sobre precio de venta).
    - Price Index = precio salida por g/ml del producto / el de FANNY del mismo SKU y zona × 100
                    (FANNY = 100; >100 = el producto es más caro que FANNY).
    - Dinámica de Sell Out = notas escritas por los vendedores en la fila FANNY.
"""
import sys
import tomllib
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BASE = Path(__file__).parent
DATA = BASE / "data"
RUTA_SECRETS = BASE / ".streamlit" / "secrets.toml"
ORDEN_ZONAS = ["MANAGUA", "SUR ORIENTE, OCCIDENTE", "NORTE, CENTRO"]
COLUMNAS = ["ZONA", "DISTRIBUIDOR", "Marca", "Producto", "Precio entrada", "Precio salida",
            "Margen %", "Margen Abs", "Price Index", "Dinamica de Sell Out", "N° PDV"]

cfg = tomllib.loads(RUTA_SECRETS.read_text(encoding="utf-8")) if RUTA_SECRETS.exists() else {}
USAR_SHEET = "gcp_service_account" in cfg and "SHEET_ID" in cfg


def leer(nombre: str) -> pd.DataFrame:
    if USAR_SHEET:
        import gspread
        libro = gspread.service_account_from_dict(cfg["gcp_service_account"]).open_by_key(cfg["SHEET_ID"])
        return pd.DataFrame(libro.worksheet(nombre).get_all_records()).astype(str)
    ruta = DATA / f"{nombre}.csv"
    return pd.read_csv(ruta, dtype=str).fillna("") if ruta.exists() else pd.DataFrame()


def arg(nombre):
    return sys.argv[sys.argv.index(nombre) + 1] if nombre in sys.argv else None


# ── 1. Datos ──
reg = leer("Registro_Precios")
if reg.empty:
    sys.exit("Registro_Precios está vacío: no hay levantamientos para reportar.")
sku, comp, rutas = leer("Maestro_SKU"), leer("Maestro_Competencia"), leer("Maestro_Ruta")

desde, hasta = arg("--desde"), arg("--hasta")
if desde:
    reg = reg[reg["Fecha"] >= desde]
if hasta:
    reg = reg[reg["Fecha"] <= hasta]
if reg.empty:
    sys.exit("No hay levantamientos en el período indicado.")

for col in ["Precio_Entrada", "Precio_Salida", "Precio_UM"]:
    reg[col] = pd.to_numeric(reg[col], errors="coerce")

# Atributos del producto (marca, distribuidor, descripción, categoría)
prod_propio = sku.rename(columns={"SKU_ID": "Producto_ID", "Proveedor": "Distribuidor"})
prod_propio["Marca"] = prod_propio.get("Marca", "FANNY")
prod_comp = comp[["Comp_ID", "Marca", "Distribuidor", "Producto"]].rename(columns={"Comp_ID": "Producto_ID"})
productos = pd.concat([prod_propio[["Producto_ID", "Marca", "Distribuidor", "Producto"]], prod_comp])
reg = (reg.merge(productos, on="Producto_ID", how="left")
          .merge(sku[["SKU_ID", "Categoria"]], on="SKU_ID", how="left"))

zona = dict(zip(rutas["Ruta"], rutas.get("Zona_Reporte", rutas["Ruta"])))
reg["ZONA"] = reg["Ruta"].map(zona).fillna(reg["Ruta"])

# ── 2. Agregación por zona y producto ──
disp = reg[reg["Disponible"] == "Sí"].copy()
# El margen solo se calcula con PDV que tienen entrada Y salida (no mezcla muestras distintas)
ambos = disp["Precio_Entrada"].notna() & disp["Precio_Salida"].notna()
disp["m_abs"] = (disp["Precio_Salida"] - disp["Precio_Entrada"]).where(ambos)
disp["s_con_e"] = disp["Precio_Salida"].where(ambos)
g = (disp.groupby(["Categoria", "ZONA", "SKU_ID", "Producto_ID", "Tipo", "Comp_Orden",
                   "Distribuidor", "Marca", "Producto"], dropna=False)
         .agg(entrada=("Precio_Entrada", "mean"), salida=("Precio_Salida", "mean"),
              pum=("Precio_UM", "mean"), n=("Cliente_Nombre", "nunique"),
              m_sum=("m_abs", "sum"), s_sum=("s_con_e", "sum"), m_n=("m_abs", "count"))
         .reset_index())
g["Margen Abs"] = (g["m_sum"] / g["m_n"]).where(g["m_n"] > 0).round(2)
g["Margen %"] = (g["m_sum"] / g["s_sum"]).where(g["m_n"] > 0).round(4)

fanny = g[g["Tipo"] == "Propio"][["ZONA", "SKU_ID", "pum"]].rename(columns={"pum": "pum_f"})
g = g.merge(fanny, on=["ZONA", "SKU_ID"], how="left")
g["Price Index"] = (g["pum"] / g["pum_f"] * 100).round(0)

notas = (reg[(reg["Tipo"] == "Propio") & (reg["Sell_Out"].fillna("") != "")]
         .groupby(["ZONA", "SKU_ID"])["Sell_Out"].agg(lambda s: "; ".join(dict.fromkeys(s))))
g["Dinamica de Sell Out"] = [notas.get((z, s), "") if t == "Propio" else ""
                             for z, s, t in zip(g["ZONA"], g["SKU_ID"], g["Tipo"])]

g["orden_z"] = g["ZONA"].map({z: i for i, z in enumerate(ORDEN_ZONAS)}).fillna(99)
g["orden_t"] = (g["Tipo"] != "Propio").astype(int)
g["Comp_Orden"] = pd.to_numeric(g["Comp_Orden"], errors="coerce")
g = g.sort_values(["Categoria", "orden_z", "SKU_ID", "orden_t", "Comp_Orden"])
g = g.rename(columns={"Distribuidor": "DISTRIBUIDOR", "entrada": "Precio entrada",
                      "salida": "Precio salida", "n": "N° PDV"})

# ── 3. Excel con el formato encomendado ──
etiqueta = f"{desde or 'inicio'}_a_{hasta or datetime.now():%Y-%m-%d}".replace("-", "")
salida = BASE / f"Reporte_Precios_{etiqueta}.xlsx"
azul = PatternFill("solid", fgColor="1F5F8B")
fanny_fill = PatternFill("solid", fgColor="E8F1F8")
borde = Border(bottom=Side(style="thin", color="D9E0E6"))

with pd.ExcelWriter(salida, engine="openpyxl") as xw:
    for cat, d in g.groupby("Categoria", sort=False):
        hoja = str(cat).upper()[:31]
        d[COLUMNAS].to_excel(xw, sheet_name=hoja, index=False, startrow=3, startcol=1)
        ws = xw.sheets[hoja]
        ws["B1"] = f"{hoja} · Levantamiento de precios SUPLISA"
        ws["B1"].font = Font(name="Arial", size=12, bold=True)
        ws["B2"] = f"Período: {desde or 'inicio'} a {hasta or 'hoy'} · Promedio de PDV levantados por zona"
        ws["B2"].font = Font(name="Arial", size=9, italic=True)
        for c in ws[4][1:]:
            c.font, c.fill = Font(name="Arial", size=10, bold=True, color="FFFFFF"), azul
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        for i, fila in enumerate(ws.iter_rows(min_row=5, max_row=4 + len(d), min_col=2, max_col=12)):
            es_fanny = d.iloc[i]["Tipo"] == "Propio"
            for c in fila:
                c.font = Font(name="Arial", size=10, bold=es_fanny)
                c.border = borde
                if es_fanny:
                    c.fill = fanny_fill
            fila[4].number_format = fila[5].number_format = fila[7].number_format = "#,##0.00"
            fila[6].number_format = "0.0%"
            fila[8].number_format = "0"
        for col, ancho in zip("BCDEFGHIJKL", [24, 14, 13, 36, 13, 13, 10, 11, 11, 40, 8]):
            ws.column_dimensions[col].width = ancho
        ws.freeze_panes = "B5"
    base = reg.drop(columns=["ZONA"]).assign(Zona=reg["ZONA"])
    base.to_excel(xw, sheet_name="Base_Levantamiento", index=False)

print(f"Período: {desde or 'inicio'} a {hasta or 'hoy'} · {reg['ID_Lev'].nunique()} levantamientos · "
      f"{reg['Cliente_Nombre'].nunique()} PDV")
print(f"Informe generado -> {salida.name}")
